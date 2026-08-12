from __future__ import annotations

import json
import os
import re
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

from PIL import Image
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


ROOT = Path(r"\\192.168.1.92\310_共享\999_Process_Control_Board\99_Continous_Improving_Program(CIP)\02_测试专题提升")
OUTPUT = Path(r"D:\CHATGPT_file\else thing\outputs\cip_file_table\inventory.json")

TEXT_EXTENSIONS = {
    ".txt", ".json", ".config", ".xml", ".ini", ".py", ".m", ".tcl",
    ".aspx", ".map", ".browser", ".info",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff"}
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".wmv"}
PROGRAM_EXTENSIONS = {".dll", ".exe", ".bin", ".assets", ".ress"}
MODEL_EXTENSIONS = {".stl", ".stlb"}


def clean_text(value: Any) -> str:
    text = str(value).replace("\x00", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def unique_chunks(chunks: list[str], limit: int = 5000) -> str:
    output: list[str] = []
    seen: set[str] = set()
    used = 0
    for chunk in chunks:
        chunk = clean_text(chunk)
        if not chunk or chunk in seen:
            continue
        seen.add(chunk)
        if used + len(chunk) > limit:
            chunk = chunk[: max(0, limit - used)]
        if chunk:
            output.append(chunk)
            used += len(chunk)
        if used >= limit:
            break
    return " | ".join(output)


def extract_xlsx(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    if size > 80 * 1024 * 1024:
        return "", {"status": "skipped_large", "reason": "工作簿超过80MB"}
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    chunks: list[str] = []
    sheet_names = wb.sheetnames
    scanned_cells = 0
    try:
        for ws in wb.worksheets[:6]:
            chunks.append(f"工作表:{ws.title}")
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 80), max_col=min(ws.max_column or 1, 20), values_only=True):
                for value in row:
                    if value is None or isinstance(value, (datetime, int, float, bool)):
                        continue
                    text = clean_text(value)
                    if len(text) >= 2:
                        chunks.append(text)
                    scanned_cells += 1
                    if scanned_cells >= 500 or sum(map(len, chunks)) >= 7000:
                        break
                if scanned_cells >= 500 or sum(map(len, chunks)) >= 7000:
                    break
            if scanned_cells >= 500 or sum(map(len, chunks)) >= 7000:
                break
    finally:
        wb.close()
    return unique_chunks(chunks), {"status": "ok", "sheets": sheet_names, "sheet_count": len(sheet_names)}


def extract_pdf(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    if size > 150 * 1024 * 1024:
        return "", {"status": "skipped_large", "reason": "PDF超过150MB"}
    reader = PdfReader(str(path), strict=False)
    encrypted = bool(reader.is_encrypted)
    if encrypted:
        result = reader.decrypt("")
        if not result:
            return "", {"status": "encrypted", "pages": len(reader.pages)}
    chunks: list[str] = []
    for page in reader.pages[:8]:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    text = unique_chunks(chunks)
    status = "ok" if text else "no_extractable_text"
    return text, {"status": status, "pages": len(reader.pages), "encrypted": encrypted}


def extract_docx(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    if size > 80 * 1024 * 1024:
        return "", {"status": "skipped_large", "reason": "Word文档超过80MB"}
    doc = Document(str(path))
    chunks: list[str] = []
    for p in doc.paragraphs[:300]:
        if p.text:
            chunks.append(p.text)
    for table in doc.tables[:10]:
        for row in table.rows[:30]:
            chunks.extend(cell.text for cell in row.cells[:10])
    return unique_chunks(chunks), {"status": "ok", "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}


def extract_pptx(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    if size > 120 * 1024 * 1024:
        return "", {"status": "skipped_large", "reason": "演示文稿超过120MB"}
    deck = Presentation(str(path))
    chunks: list[str] = []
    for slide in deck.slides[:30]:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
        if sum(map(len, chunks)) >= 7000:
            break
    return unique_chunks(chunks), {"status": "ok", "slides": len(deck.slides)}


def extract_text_file(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    raw = path.read_bytes()[:65536]
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            text = raw.decode(encoding)
            return unique_chunks([text]), {"status": "ok", "encoding": encoding, "truncated": size > len(raw)}
        except UnicodeDecodeError:
            continue
    return "", {"status": "binary_or_unknown_encoding"}


def extract_image(path: Path) -> tuple[str, dict[str, Any]]:
    with Image.open(path) as image:
        info = {
            "status": "ok",
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
            "format": image.format,
            "frames": getattr(image, "n_frames", 1),
        }
    return "", info


def detect_signature(path: Path) -> str:
    try:
        data = path.read_bytes()[:16]
    except Exception:
        return ""
    if data.startswith(b"MZ"):
        return "Windows PE程序/依赖文件"
    if data.startswith(b"PK\x03\x04"):
        return "ZIP/OpenXML压缩文件"
    if data.startswith(b"%PDF"):
        return "PDF文件"
    if data[:4] in (b"\x00\x00\x00\x18", b"\x00\x00\x00\x20") and b"ftyp" in data:
        return "MP4媒体文件"
    return ""


def extract_file(path: Path, size: int) -> tuple[str, dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return extract_xlsx(path, size)
    if ext == ".pdf":
        return extract_pdf(path, size)
    if ext == ".docx":
        return extract_docx(path, size)
    if ext == ".pptx":
        return extract_pptx(path, size)
    if ext in TEXT_EXTENSIONS:
        return extract_text_file(path, size)
    if ext in IMAGE_EXTENSIONS:
        return extract_image(path)
    if ext == ".xls":
        return "", {"status": "needs_local_app", "reason": "旧版XLS需本机Excel读取"}
    if ext in VIDEO_EXTENSIONS:
        return "", {"status": "metadata_only", "kind": "video"}
    if ext in PROGRAM_EXTENSIONS:
        return "", {"status": "metadata_only", "kind": "program"}
    if ext in MODEL_EXTENSIONS:
        return "", {"status": "metadata_only", "kind": "3d_model"}
    signature = detect_signature(path)
    return "", {"status": "metadata_only", "signature": signature}


TOPIC_MAP = {
    "00_archive": "归档",
    "01_test_specification": "测试规范",
    "02_test_equipment": "测试设备",
    "03_lab_management": "实验室管理",
    "04_sunlight-overheat": "阳光倒灌",
    "05_working_instruction": "仪器资料",
    "06_test_input": "测试输入",
    "07_eol": "EOL",
    "08_warping": "畸变校正",
    "09_stanard&norm": "标准与法规",
    "77_personal_output": "个人输出",
    "88_competencies_matrix": "能力矩阵",
    "99_review_record": "评审记录",
    "测试规范": "测试规范",
}


KEYWORDS = [
    ("畸变|warping|distortion", "畸变校正"),
    ("阳光倒灌|sunlight|overheat", "阳光倒灌"),
    ("测试规范|test specification|specification", "测试规范"),
    ("光学性能|optical performance", "光学性能测试"),
    ("操作手册|使用手册|working instruction|manual", "操作说明"),
    ("会议纪要|会议记录|meeting", "会议纪要"),
    ("采购|选型|报价|供应商", "采购选型"),
    ("实验室|lab|5s|门禁|排班", "实验室管理"),
    ("测试报告|检测报告|test report", "测试报告"),
    ("标准|gb/t|gjb|regulation|standard", "标准法规"),
    ("台架|equipment|设备", "测试设备"),
    ("双目|相机|camera", "相机/视觉设备"),
    ("伺服|电机|motor", "伺服机构"),
    ("扫描仪|3d", "3D扫描"),
    ("能力矩阵|competenc", "能力矩阵"),
    ("review|评审", "评审记录"),
]


def relative_parts(path: Path) -> list[str]:
    rel = os.path.relpath(str(path), str(ROOT))
    return rel.split(os.sep)


def infer_topic(parts: list[str], text: str) -> str:
    top = parts[0].lower() if parts else ""
    if top == "01_test_specification" and len(parts) > 1:
        second = parts[1].lower()
        if "warping" in second:
            return "畸变校正"
        if "optical" in second:
            return "光学性能测试"
    path_text = " ".join(parts)
    haystack = f"{path_text} {text[:2500]}".lower()
    for pattern, label in KEYWORDS:
        if re.search(pattern, haystack, flags=re.I):
            return label
    return TOPIC_MAP.get(top, parts[0] if parts else "其他")


def useful_excerpt(text: str, stem: str, max_len: int = 90) -> str:
    candidates = [clean_text(x) for x in re.split(r"[|\n\r。；;]", text)]
    generic = re.compile(r"^(工作表:|sheet\d*$|目录$|序号$|内容$|备注$|name$|date$)", re.I)
    for candidate in candidates:
        if len(candidate) < 4 or generic.search(candidate):
            continue
        if re.fullmatch(r"[\d\W_]+", candidate):
            continue
        return candidate[:max_len]
    return stem[:max_len]


def describe(path: Path, parts: list[str], text: str, meta: dict[str, Any], topic: str) -> str:
    ext = path.suffix.lower()
    stem = path.stem
    status = meta.get("status", "")
    excerpt = useful_excerpt(text, stem)
    context = topic if topic else (parts[-2] if len(parts) > 1 else "相关资料")

    if ext == ".xlsx":
        if status == "ok":
            sheets = "、".join(meta.get("sheets", [])[:5])
            return f"Excel工作簿，工作表：{sheets}；主要内容：{excerpt}"
        if status in {"encrypted", "error"}:
            return f"Excel文件，自动读取失败，需本机加密软件/Excel只读确认；用途：{context}"
        return f"Excel文件；用途根据目录和文件名判断为：{context}"
    if ext == ".xls":
        return f"旧版Excel文件，需本机Excel只读确认；用途：{context}"
    if ext == ".pdf":
        pages = meta.get("pages")
        if text:
            return f"PDF文档，共{pages}页；主要内容：{excerpt}"
        if status == "encrypted":
            return f"加密PDF，共{pages}页，需本机加密软件只读确认；用途：{context}"
        return f"PDF文档，共{pages or '未知'}页，未提取到可搜索文字（可能为扫描件）；用途：{context}"
    if ext == ".docx":
        if text:
            return f"Word文档；主要内容：{excerpt}"
        return f"Word文档，未提取到正文；用途：{context}"
    if ext == ".pptx":
        slides = meta.get("slides", "未知")
        if text:
            return f"PPT演示文稿，共{slides}页；主要内容：{excerpt}"
        return f"PPT演示文稿，共{slides}页，未提取到可搜索文字；用途：{context}"
    if ext in IMAGE_EXTENSIONS:
        dims = f"{meta.get('width', '?')}×{meta.get('height', '?')}像素"
        if "test_input" in " ".join(parts).lower() or "input" in " ".join(parts).lower():
            return f"测试输入/示例图片，{dims}；内容用途：{context}"
        if any(word in " ".join(parts).lower() for word in ("photo", "照片", "pic", "image")):
            return f"现场/设备参考图片，{dims}；内容用途：{context}"
        return f"图片文件，{dims}；内容用途根据目录和文件名判断为：{context}"
    if ext in VIDEO_EXTENSIONS:
        return f"演示/记录视频；内容用途根据目录和文件名判断为：{context}"
    if ext in MODEL_EXTENSIONS:
        return f"3D模型文件；用于{context}相关结构或测试对象"
    if ext in {".m", ".py", ".tcl", ".aspx"}:
        if text:
            return f"程序源代码/脚本；主要内容：{excerpt}"
        return f"程序源代码/脚本；用于{context}相关处理"
    if ext in PROGRAM_EXTENSIONS or meta.get("kind") == "program" or meta.get("signature") == "Windows PE程序/依赖文件":
        parent_text = " ".join(parts).lower()
        if "unity" in parent_text:
            return f"Unity程序运行组件/依赖库；服务于{context}工具"
        return f"程序文件或运行依赖；服务于{context}相关工具"
    if ext in TEXT_EXTENSIONS:
        if text:
            return f"文本/配置文件；主要内容：{excerpt}"
        return f"文本/配置文件；用于{context}相关程序或记录"
    signature = meta.get("signature")
    if signature:
        return f"{signature}；内容用途根据目录和文件名判断为：{context}"
    return f"{ext.lstrip('.').upper() or '无扩展名'}文件；内容用途根据目录和文件名判断为：{context}"


def file_type(path: Path, meta: dict[str, Any]) -> str:
    ext = path.suffix
    if ext:
        return ext.lstrip(".").upper()
    signature = meta.get("signature", "")
    if "程序" in signature:
        return "EXE/无扩展名"
    return "无扩展名"


def main() -> None:
    files: list[Path] = []
    walk_errors: list[dict[str, str]] = []

    def onerror(error: OSError) -> None:
        walk_errors.append({"path": getattr(error, "filename", ""), "error": str(error)})

    for current, _, names in os.walk(ROOT, onerror=onerror):
        for name in names:
            files.append(Path(current) / name)
    files.sort(key=lambda p: str(p).lower())

    records: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    for index, path in enumerate(files, start=1):
        try:
            stat = path.stat()
            size = stat.st_size
            modified = datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds")
        except Exception as exc:
            failures.append({"path": str(path), "error": f"stat: {exc}"})
            continue

        parts = relative_parts(path)
        try:
            text, meta = extract_file(path, size)
        except Exception as exc:
            text = ""
            error_name = type(exc).__name__
            error_text = clean_text(exc)
            encrypted_hint = any(token in f"{error_name} {error_text}".lower() for token in ("encrypted", "password", "file is not a zip", "badzip"))
            meta = {
                "status": "encrypted_or_protected" if encrypted_hint else "error",
                "error_type": error_name,
                "error": error_text[:500],
            }
            failures.append({"path": str(path), "error": f"{error_name}: {error_text}"})

        topic = infer_topic(parts, text)
        record = {
            "index": index,
            "group": parts[0] if parts else "",
            "directory": str(path.parent),
            "relative_path": os.path.relpath(str(path), str(ROOT)),
            "file_name": path.stem,
            "full_name": path.name,
            "extension": path.suffix,
            "file_type": file_type(path, meta),
            "topic": topic,
            "description": describe(path, parts, text, meta, topic),
            "content_excerpt": text[:1000],
            "size_bytes": size,
            "modified": modified,
            "extraction": meta,
        }
        records.append(record)

        if index % 50 == 0 or index == len(files):
            print(f"PROGRESS {index}/{len(files)}", flush=True)

    payload = {
        "root": str(ROOT),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "file_count_discovered": len(files),
        "record_count": len(records),
        "walk_errors": walk_errors,
        "failures": failures,
        "records": records,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "file_count": len(files),
        "record_count": len(records),
        "failure_count": len(failures),
        "walk_error_count": len(walk_errors),
        "output": str(OUTPUT),
    }, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
